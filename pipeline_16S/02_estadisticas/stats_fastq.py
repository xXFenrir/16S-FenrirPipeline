#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse, gzip, os, re, sys, time, shutil, subprocess as sp, fnmatch
from pathlib import Path
from typing import Optional, List

# ---------- Utilidades básicas ----------

def is_gzip(p: Path) -> bool:
    return str(p).endswith(".gz")

def open_maybe_gzip(p: Path):
    if is_gzip(p):
        return gzip.open(p, "rt", encoding="ascii", errors="ignore")
    return open(p, "rt", encoding="ascii", errors="ignore")

def iter_fastq_reads(handle):
    while True:
        h = handle.readline()
        if not h:
            break
        seq = handle.readline()
        plus = handle.readline()
        qual = handle.readline()
        if not qual:
            break
        yield seq.strip(), qual.strip()

def n50_from_lengths(lengths, total_bases):
    if not lengths or total_bases == 0:
        return 0
    ls = sorted(lengths, reverse=True)
    half = total_bases / 2
    acc = 0
    for L in ls:
        acc += L
        if acc >= half:
            return L
    return 0

# ---------- PRIMERS ----------

def read_fasta_seqs(path: Path) -> List[str]:
    seqs, cur = [], []
    with open(path, "rt", encoding="utf-8", errors="ignore") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            if line.startswith(">"):
                if cur:
                    seqs.append("".join(cur).upper())
                    cur = []
            else:
                cur.append(re.sub(r"[^ACGTNacgtn]", "", line))
        if cur:
            seqs.append("".join(cur).upper())
    return [s for s in seqs if s]

_rcmap = str.maketrans("ACGTN", "TGCAN")
def revcomp(s: str) -> str:
    return s.upper().translate(_rcmap)[::-1]

def hamming_leq_k(a: str, b: str, k: int) -> bool:
    mism = 0
    for ca, cb in zip(a, b):
        if cb == "N":  # N en el patrón no penaliza
            continue
        if ca != cb:
            mism += 1
            if mism > k:
                return False
    return True

def any_primer_in_window(seq: str, primers: List[str], k: int) -> bool:
    n = len(seq)
    for p in primers:
        m = len(p)
        if m > n:
            continue
        prc = revcomp(p)
        for i in range(0, n - m + 1):
            sub = seq[i:i+m]
            if hamming_leq_k(sub, p, k) or hamming_leq_k(sub, prc, k):
                return True
    return False

# ---------- Métricas por archivo ----------

def process_fastq(path: Path,
                  phred_offset: int = 33,
                  max_reads: Optional[int] = None,
                  primers: Optional[List[str]] = None,
                  primer_window: int = 120,
                  primer_max_mismatches: int = 2,
                  primer_scan: int = 10000) -> dict:
    reads = total_bases = gc_bases = n_bases = qsum = 0
    lengths = []
    rq_min = float("inf")
    rq_max = float("-inf")

    primer_scanned = primer_any_hits = 0
    do_primers = primers is not None and len(primers) > 0 and primer_scan > 0

    with open_maybe_gzip(path) as fh:
        for i, (seq, qual) in enumerate(iter_fastq_reads(fh), start=1):
            L = len(seq)
            if L == 0:
                continue
            lengths.append(L); reads += 1; total_bases += L

            sup = seq.upper()
            gc_bases += sup.count("G") + sup.count("C")
            n_bases  += sup.count("N")

            Lq = min(L, len(qual))
            if Lq > 0:
                r_qsum = 0
                for ch in qual[:Lq]:
                    q = (ord(ch) - phred_offset)
                    r_qsum += q
                    qsum += q
                r_qmean = r_qsum / Lq
                if r_qmean < rq_min: rq_min = r_qmean
                if r_qmean > rq_max: rq_max = r_qmean

            if do_primers and primer_scanned < primer_scan:
                head = sup[:primer_window]
                tail = sup[-primer_window:] if L >= primer_window else sup
                hit_head = any_primer_in_window(head, primers, primer_max_mismatches)
                hit_tail = any_primer_in_window(tail, primers, primer_max_mismatches) if L > primer_window else False
                if hit_head or hit_tail:
                    primer_any_hits += 1
                primer_scanned += 1

            if max_reads is not None and reads >= max_reads:
                break

    if reads == 0 or total_bases == 0:
        mean_len = min_len = max_len = n50 = 0
        gc_pct = mean_q = 0.0
        rq_min_out = rq_max_out = 0.0
    else:
        mean_len = round(total_bases / reads, 2)
        min_len = min(lengths); max_len = max(lengths)
        n50 = n50_from_lengths(lengths, total_bases)
        gc_pct = round(100.0 * gc_bases / total_bases, 2)
        mean_q = round(qsum / total_bases, 2)
        rq_min_out = 0.0 if rq_min == float("inf") else round(rq_min, 2)
        rq_max_out = 0.0 if rq_max == float("-inf") else round(rq_max, 2)

    if do_primers and primer_scanned > 0:
        primer_any_pct = round(100.0 * primer_any_hits / primer_scanned, 2)
    else:
        primer_any_pct = ""

    return {
        "sample": path.stem.replace(".fastq","").replace(".fq",""),
        "reads": reads,
        "bases": total_bases,
        "mean_len": mean_len,
        "min_len": min_len if reads and total_bases else 0,
        "max_len": max_len if reads and total_bases else 0,
        "N50": n50,
        "GC_percent": gc_pct,
        "mean_Q": mean_q,                  # promedio global por base
        "Qmean_read_min": rq_min_out,      # mínimo del promedio por lectura
        "Qmean_read_max": rq_max_out,      # máximo del promedio por lectura
        "primer_any_percent": primer_any_pct,
    }

# ---------- Construcción de la tabla (nombres en español) ----------

HEADERS = [
    "Muestra","lecturas","bases","longitud promedio","longitud mínima","longitud máxima",
    "N50","GC%","QScore promedio","QScore mínimo (por lectura)","QScore máximo (por lectura)","% primers"
]

ORDER = [
    "sample","reads","bases","mean_len","min_len","max_len",
    "N50","GC_percent","mean_Q","Qmean_read_min","Qmean_read_max","primer_any_percent"
]

def _fmt(x, dec: int = 2):
    if x == "" or x is None:
        return ""
    if isinstance(x, float):
        s = f"{x:.{dec}f}"
        s = s.rstrip('0').rstrip('.')
        return s
    return str(x)

def row_to_display(row: dict) -> list:
    mapping = {
        "sample": row["sample"],
        "reads": _fmt(row["reads"], 0),
        "bases": _fmt(row["bases"], 0),
        "mean_len": _fmt(row["mean_len"], 2),
        "min_len": _fmt(row["min_len"], 0),
        "max_len": _fmt(row["max_len"], 0),
        "N50": _fmt(row["N50"], 0),
        "GC_percent": _fmt(row["GC_percent"], 2),
        "mean_Q": _fmt(row["mean_Q"], 2),
        "Qmean_read_min": _fmt(row["Qmean_read_min"], 2),
        "Qmean_read_max": _fmt(row["Qmean_read_max"], 2),
        "primer_any_percent": _fmt(row["primer_any_percent"], 2),
    }
    return [mapping[k] for k in ORDER]

def row_to_excel(row: dict) -> dict:
    return {
        "Muestra": row["sample"],
        "lecturas": int(row["reads"]) if isinstance(row["reads"], (int,float)) else None,
        "bases": int(row["bases"]) if isinstance(row["bases"], (int,float)) else None,
        "longitud promedio": round(float(row["mean_len"]), 2) if row["mean_len"] != "" else None,
        "longitud mínima": int(row["min_len"]) if row["min_len"] != "" else None,
        "longitud máxima": int(row["max_len"]) if row["max_len"] != "" else None,
        "N50": int(row["N50"]) if row["N50"] != "" else None,
        "GC%": round(float(row["GC_percent"]), 2) if row["GC_percent"] != "" else None,
        "QScore promedio": round(float(row["mean_Q"]), 2) if row["mean_Q"] != "" else None,
        "QScore mínimo (por lectura)": round(float(row["Qmean_read_min"]), 2) if row["Qmean_read_min"] != "" else None,
        "QScore máximo (por lectura)": round(float(row["Qmean_read_max"]), 2) if row["Qmean_read_max"] != "" else None,
        "% primers": round(float(row["primer_any_percent"]), 2) if row["primer_any_percent"] != "" else None,
    }

# ---------- Salidas: TXT bonito + TSV + XLSX ----------

def write_pretty_table(rows: list, headers: list, out_path: Path):
    str_rows = [row_to_display(r) for r in rows]
    widths = [len(h) for h in headers]
    for r in str_rows:
        for i, cell in enumerate(r):
            widths[i] = max(widths[i], len(str(cell)))
    def mkline(cells):
        parts = [(str(cells[i]).ljust(widths[i])) for i in range(len(headers))]
        return "| " + " | ".join(parts) + " |"
    header_line = mkline(headers)
    sep_line = "-" * len(header_line)
    with open(out_path, "w", encoding="utf-8") as out:
        out.write(header_line + "\n")
        out.write(sep_line + "\n")
        for r in str_rows:
            out.write(mkline(r) + "\n")
            out.write(sep_line + "\n")

def write_tsv(rows: list, out_path: Path):
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\t".join(HEADERS) + "\n")
        for r in rows:
            ex = row_to_excel(r)
            f.write("\t".join("" if ex[h] is None else str(ex[h]) for h in HEADERS) + "\n")

def try_write_xlsx(rows: list, out_xlsx: Path, tmp_tsv: Optional[Path]=None) -> bool:
    # 1) pandas
    try:
        import pandas as pd
        df = pd.DataFrame([row_to_excel(r) for r in rows], columns=HEADERS)
        df.to_excel(out_xlsx, index=False)
        print("[INFO] XLSX escrito con pandas", file=sys.stderr)
        return True
    except Exception as e:
        print(f"[INFO] pandas no disponible/usable ({e.__class__.__name__}): intento openpyxl...", file=sys.stderr)
    # 2) openpyxl
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        for r in rows:
            ex = row_to_excel(r)
            ws.append([ex[h] for h in HEADERS])
        wb.save(out_xlsx)
        print("[INFO] XLSX escrito con openpyxl", file=sys.stderr)
        return True
    except Exception as e:
        print(f"[INFO] openpyxl no disponible/usable ({e.__class__.__name__}): intento LibreOffice...", file=sys.stderr)
    # 3) LibreOffice headless (fallback)
    soffice = shutil.which("libreoffice") or shutil.which("soffice")
    if soffice and tmp_tsv:
        write_tsv(rows, tmp_tsv)
        try:
            cmd = [
                soffice, "--headless",
                "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                "--infilter=Text - txt - csv (StarCalc):9,34,76,1",
                str(tmp_tsv), "--outdir", str(out_xlsx.parent)
            ]
            sp.run(cmd, check=True, stdout=sp.PIPE, stderr=sp.PIPE)
            generated = out_xlsx.parent / (tmp_tsv.stem + ".xlsx")
            if generated.exists() and generated != out_xlsx:
                generated.replace(out_xlsx)
            print("[INFO] XLSX convertido con LibreOffice (TAB forzado)", file=sys.stderr)
            return True
        except Exception as e:
            print(f"[INFO] Filtro largo falló ({e.__class__.__name__}), probando CSV:9,34,76,1...", file=sys.stderr)
            try:
                cmd = [
                    soffice, "--headless",
                    "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                    "--infilter=CSV:9,34,76,1",
                    str(tmp_tsv), "--outdir", str(out_xlsx.parent)
                ]
                sp.run(cmd, check=True, stdout=sp.PIPE, stderr=sp.PIPE)
                generated = out_xlsx.parent / (tmp_tsv.stem + ".xlsx")
                if generated.exists() and generated != out_xlsx:
                    generated.replace(out_xlsx)
                print("[INFO] XLSX convertido con LibreOffice (CSV:9,34,76,1)", file=sys.stderr)
                return True
            except Exception:
                return False
    return False

def format_xlsx_two_decimals(xlsx_path: Path):
    """Fija formato '0.00' en columnas decimales del XLSX."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("[INFO] No se pudo fijar formato en XLSX: falta openpyxl (instala con 'pip install openpyxl')", file=sys.stderr)
        return

    wb = load_workbook(xlsx_path)
    ws = wb.active

    decimal_headers = {
        "longitud promedio", "GC%", "QScore promedio",
        "QScore mínimo (por lectura)", "QScore máximo (por lectura)", "% primers"
    }
    headers = [cell.value for cell in ws[1]]

    for col_idx, header in enumerate(headers, start=1):
        if header in decimal_headers:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col_idx)
                if cell.value is None:
                    continue
                try:
                    cell.value = float(cell.value)
                except Exception:
                    continue
                cell.number_format = "0.00"

    wb.save(xlsx_path)
    print("[INFO] Formato XLSX fijado a 2 decimales", file=sys.stderr)

# ---------- Descubrimiento, CLI y main ----------

def find_fastqs(root: Path, recursive: bool = False):
    files = []
    if root.is_file():
        if any(str(root).endswith(ext) for ext in (".fastq",".fq",".fastq.gz",".fq.gz")):
            files.append(root)
    else:
        pats = ["*.fastq","*.fq","*.fastq.gz","*.fq.gz"]
        if recursive:
            for pat in pats: files.extend(root.rglob(pat))
        else:
            for pat in pats: files.extend(root.glob(pat))
    return sorted(set(files))

def parse_args():
    p = argparse.ArgumentParser(description="Métricas FASTQ; genera TXT 'bonito' y XLSX. Logs del método usado por stderr.")
    p.add_argument("input", help="Carpeta o archivo FASTQ/FASTQ.GZ")
    p.add_argument("-o","--output", default="estadisticas_pretty.txt", help="Ruta de salida .txt (tabla)")
    p.add_argument("--xlsx-out", default=None, help="Ruta de salida .xlsx (por defecto mismo nombre que -o)")
    p.add_argument("-r","--recursive", action="store_true", help="Buscar recursivamente en subcarpetas")
    p.add_argument("--phred", type=int, default=33, help="Offset Phred (por defecto: 33)")
    p.add_argument("--max-reads", type=int, default=None, help="Máximo de lecturas por archivo (pruebas)")
    p.add_argument("--primers", type=str, default=None, help="FASTA con primers (A/C/G/T/N)")
    p.add_argument("--primer-window", type=int, default=120, help="Ventana en extremos para buscar primers [120]")
    p.add_argument("--primer-max-mismatches", type=int, default=2, help="Máximo de mismatches [2]")
    p.add_argument("--primer-scan", type=int, default=10000, help="N lecturas a escanear para primers [10000]")
    p.add_argument("--name-pattern", type=str, default="*_clean.fastq*",
                   help="Patrón (glob) para filtrar nombres cuando input es carpeta. Ej: '*_clean.fastq*', 'SRR123*.fastq.gz'")
    return p.parse_args()

def main():
    args = parse_args()
    inp = Path(args.input).expanduser()
    out_txt = Path(args.output)
    out_xlsx = Path(args.xlsx_out) if args.xlsx_out else out_txt.with_suffix(".xlsx")
    tmp_tsv = out_txt.with_suffix(".tsv")

    t0 = time.time()
    print(f"[INFO] Iniciando: analizando FASTQ en {inp} ...", file=sys.stderr)

    primers = None
    if args.primers:
        pp = Path(args.primers).expanduser()
        if pp.exists():
            primers = read_fasta_seqs(pp)
            if not primers:
                print(f"[WARN] El archivo de primers está vacío o no legible: {pp}", file=sys.stderr)
                primers = None
        else:
            print(f"[WARN] Archivo de primers no existe: {pp}. Continuaré sin primers.", file=sys.stderr)

    files = find_fastqs(inp, recursive=args.recursive)
    if not files:
        print(f"[WARN] No se encontraron archivos FASTQ en: {inp}", file=sys.stderr)
        return

    # Filtro por patrón solo cuando 'input' es carpeta
    if inp.is_dir():
        before = len(files)
        files = [f for f in files if fnmatch.fnmatch(f.name, args.name_pattern)]
        print(f"[INFO] Filtro name-pattern='{args.name_pattern}': {before} → {len(files)} archivos", file=sys.stderr)
        if not files:
            print(f"[WARN] El patrón no coincidió con ningún archivo en {inp}.", file=sys.stderr)
            return
    else:
        print(f"[INFO] 'input' es archivo; se ignora --name-pattern.", file=sys.stderr)

    print(f"[INFO] Archivos a procesar: {len(files)}", file=sys.stderr)

    rows = []
    for fpath in files:
        rows.append(process_fastq(
            fpath,
            phred_offset=args.phred,
            max_reads=args.max_reads,
            primers=primers,
            primer_window=args.primer_window,
            primer_max_mismatches=args.primer_max_mismatches,
            primer_scan=args.primer_scan
        ))

    # TXT bonito
    write_pretty_table(rows, HEADERS, out_txt)
    print(f"[OK] TXT guardado en: {out_txt.resolve()}", file=sys.stderr)

    # XLSX
    if try_write_xlsx(rows, out_xlsx, tmp_tsv):
        format_xlsx_two_decimals(out_xlsx)
        print(f"[OK] XLSX guardado en: {out_xlsx.resolve()}", file=sys.stderr)
        if tmp_tsv.exists():
            try: tmp_tsv.unlink()
            except Exception: pass
    else:
        write_tsv(rows, tmp_tsv)
        print(f"[WARN] No se pudo crear .xlsx automáticamente.", file=sys.stderr)
        print(f"       TSV para Excel: {tmp_tsv.resolve()}", file=sys.stderr)
        print(f"       Opciones:", file=sys.stderr)
        print(f"         - Instalar:  python3 -m pip install pandas openpyxl", file=sys.stderr)
        print(f"         - O convertir con LibreOffice:", file=sys.stderr)
        print(f"           libreoffice --headless --convert-to xlsx \"{tmp_tsv}\" --outdir \"{out_xlsx.parent}\"", file=sys.stderr)

    dt = time.time() - t0
    print(f"[OK] Proceso completo. Tiempo: {dt:.1f}s", file=sys.stderr)

if __name__ == "__main__":
    main()

